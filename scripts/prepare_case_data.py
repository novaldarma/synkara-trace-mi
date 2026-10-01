"""Prepare the Case 2 Excel sources for inspection and later Supabase import.

Run from any directory: python scripts/prepare_case_data.py
Requirement: python -m pip install openpyxl

Input (relative to the project root by default):
    imports/raw/case2/Incident Database/Incident Database.xlsx
    imports/raw/case2/Production Data/Production Data - RCA*.xlsx
    imports/raw/case2/Equipment Performance/Equipment Performance - RCA*.xlsx

Output: imports/prepared/case2/ (JSONL, source_catalog.json, manifest.json).
This script does not read RCA PowerPoint or external energy data; those have
separate preparation steps. It never connects to Supabase or an AI provider.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import sys
import tempfile
import uuid
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime, time, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "openpyxl belum terpasang. Jalankan: python -m pip install openpyxl"
    ) from exc


DATASET_ID = "caliber2026_case2"
SCHEMA_VERSION = 1
ASSETS = {
    1: ("PU-2101B", "ARP"),
    2: ("KO-3201", "ZCU"),
    3: ("PM-4405B", "NUP"),
    4: ("HE-3301", "ZCU"),
    5: ("BL-5702", "OPP"),
}
INCIDENT_HEADERS = (
    "Serial No", "MTO No.", "AR No.", "Plant", "Tag Number", "Eq. Class",
    "Date of Occur.", "Risk Case Title", "Highest Impact", "Pre-Risk",
    "Risk Score", "PIC (RCA)", "Overall Status", "Discipline", "Eq. Type",
    "Component", "F Mechanism", "Downtime (hrs)", "Act. Loss (k US$)",
    "Pot. Loss (k US$)", "Total Loss (k US$)", "RCA Due Date",
    "Month - Year",
)
PI_HEADERS = (
    "Name", "Description", "digitalset", "engunits", "span",
    "typicalvalue", "zero", "instrumenttag",
)
SUMMARY_HEADERS = ("KPI", "Value", "Basis / Formula")
SUMMARY_KPIS = {
    "Monitoring Period (weeks)", "Total Downtime (hours)", "Period Hours",
    "Availability (%)", "No. of Failures (period)", "MTBF (hours)",
    "MTTR (hours)", "ALARM readings", "TRIP readings", "NORMAL readings",
    "PM Compliance (%)", "Production Loss (ton)", "Estimated Loss (k USD)",
}
INFO_LABELS = {
    "Equipment Tag", "Equipment Name", "Equipment Type", "Equipment Class",
    "Plant / Unit", "Discipline", "Criticality", "Design Life",
    "Monitoring Method", "Linked RCA / AR No.", "Failure Date",
    "Dominant Failure Mode",
}
POST_EVENT_INFO = {"Linked RCA / AR No.", "Failure Date", "Dominant Failure Mode"}
OUTPUTS = (
    "incidents.jsonl", "incident_dashboard_rows.jsonl", "production_tags.jsonl",
    "production_readings.jsonl", "equipment_info_rows.jsonl",
    "equipment_metadata.jsonl", "equipment_limits.jsonl",
    "equipment_conditions.jsonl", "equipment_summary.jsonl",
)
EXPECTED = {
    "incidents.jsonl": 380,
    "production_tags.jsonl": 35,
    "production_readings.jsonl": 3600,
    "equipment_metadata.jsonl": 60,
    "equipment_limits.jsonl": 20,
    "equipment_conditions.jsonl": 130,
    "equipment_summary.jsonl": 65,
}
SENTINEL = "PREPARED_BY_TRACE_MI.txt"


class PreparationError(Exception):
    """The sources differ from the approved Case 2 inventory."""


def fail(message: str) -> None:
    raise PreparationError(message)


def digest_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_value(value: Any) -> Any:
    """Serialize cell values without inventing a timezone or changing blanks."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    fail(f"Tipe sel Excel tidak didukung: {type(value).__name__}")


def rows_with_numbers(sheet: Any) -> Iterable[tuple[int, tuple[Any, ...]]]:
    for row_number, cells in enumerate(sheet.iter_rows(), start=1):
        if any(cell.value is not None for cell in cells):
            yield row_number, cells


def cell_payload(headers: Iterable[str], cells: Iterable[Any]) -> dict[str, Any]:
    names = tuple(headers)
    items = tuple(cells)
    if len(names) != len(items):
        fail(f"Panjang header/sel tidak cocok: {len(names)} vs {len(items)}")
    raw = {name: json_value(cell.value) for name, cell in zip(names, items)}
    types = {name: getattr(cell, "data_type", "n") for name, cell in zip(names, items)}
    formats = {
        name: cell.number_format
        for name, cell in zip(names, items)
        if getattr(cell, "number_format", "General") != "General"
    }
    return {"raw": raw, "raw_types": types, "raw_number_formats": formats}


def get_header(sheet: Any, row_number: int) -> tuple[Any, ...]:
    return tuple(cell.value for cell in next(
        sheet.iter_rows(min_row=row_number, max_row=row_number)
    ))


def check_header(sheet: Any, row_number: int, expected: tuple[str, ...]) -> None:
    actual = get_header(sheet, row_number)
    if actual != expected:
        fail(
            f"Header berubah: {sheet.title} baris {row_number}. "
            f"Diharapkan {expected!r}; ditemukan {actual!r}"
        )


def iso_date(value: Any, location: str) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str):
        fail(f"Tanggal kosong/tidak dikenal pada {location}: {value!r}")
    try:
        return date.fromisoformat(value.strip()).isoformat()
    except ValueError as exc:
        raise PreparationError(f"Tanggal tidak valid pada {location}: {value!r}") from exc


def iso_timestamp(value: Any, location: str) -> str:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.strip())
        except ValueError as exc:
            raise PreparationError(
                f"Timestamp tidak valid pada {location}: {value!r}"
            ) from exc
    else:
        fail(f"Timestamp kosong/tidak dikenal pada {location}: {value!r}")
    if parsed.tzinfo is not None:
        fail(f"Zona waktu tak terduga pada {location}; tinjau sumber sebelum impor")
    return parsed.isoformat(timespec="seconds")  # naive: source gives no timezone


def discover(raw_dir: Path) -> dict[str, Any]:
    incident = raw_dir / "Incident Database" / "Incident Database.xlsx"
    if not incident.is_file():
        fail(f"File insiden tidak ditemukan: {incident}")
    result: dict[str, Any] = {"incident": incident}
    for category, folder in (
        ("production", "Production Data"),
        ("equipment", "Equipment Performance"),
    ):
        directory = raw_dir / folder
        if not directory.is_dir():
            fail(f"Folder sumber tidak ditemukan: {directory}")
        pattern = re.compile(
            rf"^{re.escape(folder)} - RCA([1-5]) (.+)\.xlsx$", re.IGNORECASE
        )
        found: dict[int, Path] = {}
        for path in directory.glob("*.xlsx"):
            if path.name.startswith("~$"):
                continue  # temporary Excel lock file, not a source
            match = pattern.fullmatch(path.name)
            if not match:
                fail(f"File XLSX di luar inventaris pada {directory}: {path.name}")
            number = int(match.group(1))
            if number in found or match.group(2).upper() != ASSETS[number][0]:
                fail(f"Nomor RCA atau tag aset tidak sesuai pada: {path.name}")
            found[number] = path
        if set(found) != set(ASSETS):
            fail(
                f"{folder}: diperlukan RCA1–RCA5; tersedia {sorted(found)}. "
                "Periksa unduhan dan ejaan nama file."
            )
        result[category] = found
    return result


class JsonlSink:
    def __init__(self, directory: Path):
        self.handles = {
            name: (directory / name).open("w", encoding="utf-8", newline="\n")
            for name in OUTPUTS
        }
        self.counts: Counter[str] = Counter()

    def write(self, filename: str, item: dict[str, Any]) -> None:
        self.handles[filename].write(json.dumps(item, ensure_ascii=False, allow_nan=False))
        self.handles[filename].write("\n")
        self.counts[filename] += 1

    def close(self) -> None:
        for handle in self.handles.values():
            handle.close()


def provenance(source: dict[str, Any], sheet: str, row: int,
               record_kind: str) -> dict[str, Any]:
    identity = f"{source['source_id']}|{sheet}|{row}"
    source_row_id = "case2_row_" + hashlib.sha256(identity.encode("utf-8")).hexdigest()
    return {
        "record_id": "case2_" + hashlib.sha256(
            f"{identity}|{record_kind}".encode("utf-8")
        ).hexdigest(),
        "source_row_id": source_row_id,
        "dataset_id": DATASET_ID,
        "source_id": source["source_id"],
        "source_file": source["relative_path"],
        "source_sha256": source["sha256"],
        "source_sheet": sheet,
        "source_row": row,  # Excel row number, one-based
    }


def source_entry(raw_dir: Path, path: Path, category: str, workbook: Any,
                 asset_tag: str | None = None, plant: str | None = None) -> dict[str, Any]:
    relative_path = path.relative_to(raw_dir).as_posix()
    sha256 = digest_file(path)
    source_id = "case2_source_" + hashlib.sha256(
        f"{DATASET_ID}|{relative_path}|{sha256}".encode("utf-8")
    ).hexdigest()
    return {
        "source_id": source_id,
        "dataset_id": DATASET_ID,
        "category": category,
        "relative_path": relative_path,
        "sha256": sha256,
        "bytes": path.stat().st_size,
        "sheets": [
            {"name": sheet.title, "excel_max_row": sheet.max_row,
             "excel_max_column": sheet.max_column}
            for sheet in workbook.worksheets
        ],
        "asset_tag": asset_tag,
        "plant": plant,
    }


def require_sheets(workbook: Any, path: Path, required: set[str]) -> None:
    if set(workbook.sheetnames) != required:
        fail(
            f"Sheet tidak sesuai pada {path.name}: "
            f"diperlukan {sorted(required)}, ada {workbook.sheetnames}"
        )


@contextmanager
def open_workbook(path: Path) -> Iterable[Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        yield workbook
    finally:
        workbook.close()


def base_record(source: dict[str, Any], sheet: str, row_number: int,
                headers: Iterable[str], cells: Iterable[Any],
                record_kind: str) -> dict[str, Any]:
    return {
        **provenance(source, sheet, row_number, record_kind),
        **cell_payload(headers, cells),
    }


def read_incident(path: Path, raw_dir: Path, sink: JsonlSink,
                  catalog: list[dict[str, Any]]) -> None:
    with open_workbook(path) as workbook:
        require_sheets(workbook, path, {"Incident Database", "Dashboard"})
        source = source_entry(raw_dir, path, "incident", workbook)
        catalog.append(source)
        sheet = workbook["Incident Database"]
        check_header(sheet, 3, INCIDENT_HEADERS)
        for row_number, cells in rows_with_numbers(sheet):
            if row_number < 4:
                continue
            if len(cells) != len(INCIDENT_HEADERS):
                fail(f"Jumlah kolom Incident berubah pada baris {row_number}")
            record = base_record(source, sheet.title, row_number, INCIDENT_HEADERS,
                                 cells, "incident")
            raw = record["raw"]
            if not isinstance(raw["Serial No"], int):
                fail(f"Serial No harus angka pada Incident baris {row_number}")
            record.update({
                "source_type": "incident",
                "occurred_date_raw": raw["Date of Occur."],
                "occurred_date": iso_date(raw["Date of Occur."], f"Incident baris {row_number}"),
                "plant": raw["Plant"],
                "asset_tag": raw["Tag Number"],
                "ar_no": None if str(raw["AR No."]).strip().lower() == "n/a"
                else raw["AR No."],
                "mto_no": raw["MTO No."],
                "rca_due_date": None if str(raw["RCA Due Date"]).strip().lower() == "n/a"
                else iso_date(raw["RCA Due Date"], f"RCA Due Date baris {row_number}"),
                "downtime_hours": raw["Downtime (hrs)"],
                "actual_loss_kusd": raw["Act. Loss (k US$)"],
                "potential_loss_kusd": raw["Pot. Loss (k US$)"],
                "total_loss_kusd_as_provided": raw["Total Loss (k US$)"],
                "overall_status_as_provided": raw["Overall Status"],
            })
            sink.write("incidents.jsonl", record)

        # The workbook's Dashboard is a supplied summary, not 380 more incidents.
        dashboard = workbook["Dashboard"]
        if dashboard.max_column != 6:
            fail(f"Dashboard columns berubah pada {path.name}")
        for row_number, cells in rows_with_numbers(dashboard):
            headers = tuple("ABCDEF")
            record = base_record(source, dashboard.title, row_number, headers,
                                 cells, "incident_dashboard")
            record["source_type"] = "incident_dashboard_source_row"
            sink.write("incident_dashboard_rows.jsonl", record)


def read_production(number: int, path: Path, raw_dir: Path, sink: JsonlSink,
                    catalog: list[dict[str, Any]]) -> None:
    asset_tag, plant = ASSETS[number]
    with open_workbook(path) as workbook:
        require_sheets(workbook, path, {"PI Tag", "Sheet2"})
        source = source_entry(raw_dir, path, "production", workbook, asset_tag, plant)
        catalog.append(source)
        tag_sheet = workbook["PI Tag"]
        check_header(tag_sheet, 1, PI_HEADERS)
        tags: dict[str, dict[str, Any]] = {}
        for row_number, cells in rows_with_numbers(tag_sheet):
            if row_number == 1:
                continue
            record = base_record(source, tag_sheet.title, row_number, PI_HEADERS,
                                 cells, "production_tag")
            raw = record["raw"]
            name = raw["Name"]
            if not isinstance(name, str) or name in tags:
                fail(f"PI Tag kosong/duplikat pada {path.name}, baris {row_number}")
            tags[name] = raw
            record.update({
                "source_type": "production_tag_metadata",
                "asset_tag": asset_tag, "plant": plant,
                "tag_name": name, "engineering_unit_as_provided": raw["engunits"],
            })
            sink.write("production_tags.jsonl", record)
        if len(tags) != 7:
            fail(f"Jumlah PI Tag pada {path.name} harus 7, ditemukan {len(tags)}")

        sheet = workbook["Sheet2"]
        header = get_header(sheet, 1)
        expected = ("Timestamp", *tags)
        if header != expected or header[-2:] != ("PLANT_RATE", "RUN_STATUS"):
            fail(f"Sheet2 dan urutan PI Tag tidak cocok pada {path.name}: {header!r}")
        timestamps: set[str] = set()
        for row_number, cells in rows_with_numbers(sheet):
            if row_number == 1:
                continue
            record = base_record(source, sheet.title, row_number, header, cells,
                                 "production_reading")
            raw = record["raw"]
            stamp = iso_timestamp(raw["Timestamp"], f"{path.name}:Sheet2:{row_number}")
            if stamp in timestamps:
                fail(f"Timestamp Production berulang pada {path.name}: {stamp}")
            timestamps.add(stamp)
            record.update({
                "source_type": "production_hourly_observation",
                "asset_tag": asset_tag, "plant": plant,
                "observed_at_raw": raw["Timestamp"],
                "observed_at_naive": stamp,
                "time_grain": "hourly_as_provided",
                "signals": [
                    {"tag_name": name, "value": raw[name],
                     "engineering_unit_as_provided": tags[name]["engunits"],
                     "instrument_tag_as_provided": tags[name]["instrumenttag"]}
                    for name in tags
                ],
            })
            sink.write("production_readings.jsonl", record)
        if len(timestamps) != 720:
            fail(f"Sheet2 {path.name}: diperlukan 720 observasi, ada {len(timestamps)}")


def parameter_and_unit(text: str) -> tuple[str, str]:
    match = re.fullmatch(r"(.+?)\s*\(([^()]+)\)", text.strip(), re.DOTALL)
    if not match:
        fail(f"Nama parameter/satuan tidak dapat dipisah: {text!r}")
    return match.group(1).strip(), match.group(2).strip()


def read_equipment(number: int, path: Path, raw_dir: Path, sink: JsonlSink,
                   catalog: list[dict[str, Any]]) -> None:
    asset_tag, plant = ASSETS[number]
    with open_workbook(path) as workbook:
        require_sheets(workbook, path, {
            "Equipment Info", "Condition History", "Performance Summary",
        })
        source = source_entry(raw_dir, path, "equipment", workbook, asset_tag, plant)
        catalog.append(source)

        info = workbook["Equipment Info"]
        if get_header(info, 4) != (
            "Equipment Tag", asset_tag, "Parameter", "Alarm / Trip"
        ):
            fail(f"Equipment Info baris 4 tidak cocok pada {path.name}")
        labels_seen: set[str] = set()
        limit_count = 0
        for row_number, cells in rows_with_numbers(info):
            if len(cells) != 4:
                fail(f"Equipment Info harus empat kolom pada {path.name}:{row_number}")
            row = base_record(source, info.title, row_number, tuple("ABCD"),
                              cells, "equipment_info_row")
            row.update({"source_type": "equipment_info_source_row",
                        "asset_tag": asset_tag, "plant": plant})
            sink.write("equipment_info_rows.jsonl", row)
            if row_number < 4:
                continue
            label = cells[0].value
            if label not in INFO_LABELS or label in labels_seen:
                fail(f"Field identitas Equipment tidak dikenal/duplikat: {path.name}:{row_number} {label!r}")
            labels_seen.add(label)
            metadata = {
                **provenance(source, info.title, row_number, "equipment_metadata"),
                "source_type": "equipment_metadata",
                "asset_tag": asset_tag, "plant": plant,
                "field_name": label,
                "value_as_provided": json_value(cells[1].value),
                "value_type_as_provided": getattr(cells[1], "data_type", "n"),
                "post_event_context": label in POST_EVENT_INFO,
            }
            sink.write("equipment_metadata.jsonl", metadata)
            if label == "Plant / Unit" and f"({plant})" not in str(cells[1].value):
                fail(f"Plant Equipment berbeda dari inventaris pada {path.name}")
            if row_number >= 5 and (cells[2].value is not None or cells[3].value is not None):
                label_raw = cells[2].value
                limits_raw = cells[3].value
                if not isinstance(label_raw, str) or not isinstance(limits_raw, str):
                    fail(f"Batas parameter tidak lengkap pada {path.name}:{row_number}")
                parameter, unit = parameter_and_unit(label_raw)
                match = re.fullmatch(r"\s*([0-9.]+)\s*/\s*([0-9.]+)\s*", limits_raw)
                if match is None:
                    fail(f"Alarm/Trip tidak dikenali pada {path.name}:{row_number}: {limits_raw!r}")
                sink.write("equipment_limits.jsonl", {
                    **provenance(source, info.title, row_number, "equipment_limit"),
                    "source_type": "equipment_threshold_reference",
                    "asset_tag": asset_tag, "plant": plant,
                    "parameter_name": parameter, "parameter_label_raw": label_raw,
                    "unit_as_provided": unit, "alarm_trip_raw": limits_raw,
                    "alarm_value_as_provided": float(match.group(1)),
                    "trip_value_as_provided": float(match.group(2)),
                    "direction": None, "historical_valid_from": None,
                    "note": "Reference values only; direction/period of validity unverified",
                })
                limit_count += 1
        if labels_seen != INFO_LABELS or limit_count != 4:
            fail(f"Equipment Info {path.name}: fields/empat batas tidak lengkap")

        condition = workbook["Condition History"]
        header = get_header(condition, 1)
        if len(header) != 8 or header[:2] != ("Week", "Date") or header[-2:] != (
            "Health Status", "Remark"
        ):
            fail(f"Header Condition History tidak cocok: {path.name}: {header!r}")
        parameters = []
        for column in header[2:6]:
            if not isinstance(column, str):
                fail(f"Nama parameter kosong pada {path.name}")
            name, unit = parameter_and_unit(column)
            parameters.append((column, name, unit))
        condition_count = 0
        for row_number, cells in rows_with_numbers(condition):
            if row_number == 1:
                continue
            record = base_record(source, condition.title, row_number, header,
                                 cells, "equipment_condition")
            raw = record["raw"]
            record.update({
                "source_type": "equipment_weekly_observation",
                "asset_tag": asset_tag, "plant": plant,
                "observed_date_raw": raw["Date"],
                "observed_date": iso_date(raw["Date"], f"{path.name}:{row_number}"),
                "observed_time": None, "available_at": None,
                "week_as_provided": raw["Week"],
                "health_status_as_provided": raw["Health Status"],
                "time_grain": "weekly_as_provided",
                "measurements": [
                    {"parameter_label_raw": column, "parameter_name": name,
                     "unit_as_provided": unit, "value": raw[column]}
                    for column, name, unit in parameters
                ],
            })
            sink.write("equipment_conditions.jsonl", record)
            condition_count += 1
        if condition_count != 26:
            fail(f"Condition History {path.name}: perlu 26 baris, ada {condition_count}")

        summary = workbook["Performance Summary"]
        check_header(summary, 3, SUMMARY_HEADERS)
        kpis: set[str] = set()
        for row_number, cells in rows_with_numbers(summary):
            if row_number <= 3:
                continue
            record = base_record(source, summary.title, row_number, SUMMARY_HEADERS,
                                 cells, "equipment_summary")
            raw = record["raw"]
            name = raw["KPI"]
            if not isinstance(name, str) or name in kpis:
                fail(f"KPI Equipment kosong/duplikat pada {path.name}:{row_number}")
            kpis.add(name)
            record.update({
                "source_type": "equipment_kpi_as_provided",
                "asset_tag": asset_tag, "plant": plant,
                "kpi_name": name, "value_as_provided": raw["Value"],
                "basis_as_provided": raw["Basis / Formula"],
                "time_scope": "full_period_retrospective_only",
            })
            sink.write("equipment_summary.jsonl", record)
        if kpis != SUMMARY_KPIS:
            fail(f"KPI Performance Summary berbeda pada {path.name}: {sorted(kpis ^ SUMMARY_KPIS)}")


def write_json(path: Path, item: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(item, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def publish(stage: Path, output_dir: Path) -> None:
    """Swap only a directory previously generated by this script."""
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    backup: Path | None = None
    allowed = set(OUTPUTS) | {"source_catalog.json", "manifest.json", SENTINEL}
    if output_dir.exists():
        if not output_dir.is_dir() or not (output_dir / SENTINEL).is_file():
            fail(f"Output sudah ada dan bukan milik script ini: {output_dir}")
        unexpected = {p.name for p in output_dir.iterdir()} - allowed
        if unexpected:
            fail(f"Folder output memuat berkas pengguna; tidak ditimpa: {sorted(unexpected)}")
        backup = output_dir.with_name(output_dir.name + ".backup-" + uuid.uuid4().hex)
        output_dir.rename(backup)
    try:
        stage.rename(output_dir)
    except OSError:
        if backup is not None:
            backup.rename(output_dir)
        raise
    if backup is not None:
        shutil.rmtree(backup)


def prepare(raw_dir: Path, output_dir: Path) -> dict[str, Any]:
    raw_dir = raw_dir.resolve()
    output_dir = output_dir.resolve()
    if not raw_dir.is_dir():
        fail(f"Folder input belum ada: {raw_dir}")
    if raw_dir == output_dir or raw_dir in output_dir.parents or output_dir in raw_dir.parents:
        fail("Folder input dan output tidak boleh berada di dalam satu sama lain")
    files = discover(raw_dir)
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=".case2-prepare-", dir=output_dir.parent
    ) as temporary:
        stage = Path(temporary) / output_dir.name
        stage.mkdir()
        catalog: list[dict[str, Any]] = []
        sink = JsonlSink(stage)
        try:
            read_incident(files["incident"], raw_dir, sink, catalog)
            for number in ASSETS:
                read_production(number, files["production"][number], raw_dir, sink, catalog)
                read_equipment(number, files["equipment"][number], raw_dir, sink, catalog)
        finally:
            sink.close()
        for name, count in EXPECTED.items():
            if sink.counts[name] != count:
                fail(f"{name}: dibutuhkan {count} catatan, ada {sink.counts[name]}")
        if len(catalog) != 11:
            fail(f"Katalog harus 11 workbook, ditemukan {len(catalog)}")
        catalog_document = {
            "dataset_id": DATASET_ID, "schema_version": SCHEMA_VERSION,
            "sources": catalog,
        }
        write_json(stage / "source_catalog.json", catalog_document)
        prepared_at = datetime.now(timezone.utc).isoformat()
        manifest = {
            "dataset_id": DATASET_ID,
            "schema_version": SCHEMA_VERSION,
            "prepared_at_utc": prepared_at,
            "source_file_count": len(catalog),
            "record_counts": {name: sink.counts[name] for name in OUTPUTS},
            "source_catalog_sha256": digest_file(stage / "source_catalog.json"),
            "output_sha256": {name: digest_file(stage / name) for name in OUTPUTS},
            "time_policy": "Keep naive Case 2 source time; no timezone invented",
            "not_included": ["RCA PowerPoint", "external UCI energy", "Supabase import"],
        }
        write_json(stage / "manifest.json", manifest)
        (stage / SENTINEL).write_text(
            "Generated by scripts/prepare_case_data.py. Edit raw sources, not these files.\n",
            encoding="utf-8",
        )
        publish(stage, output_dir)
        return manifest


def main() -> int:
    project_root = Path(__file__).resolve().parent.parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--raw-dir", type=Path, default=project_root / "imports" / "raw" / "case2",
        help="Folder Case 2 berisi subfolder Incident Database, Production Data, Equipment Performance",
    )
    parser.add_argument(
        "--output-dir", type=Path,
        default=project_root / "imports" / "prepared" / "case2",
        help="Folder keluaran terstruktur; berada di luar folder input",
    )
    args = parser.parse_args()
    try:
        manifest = prepare(args.raw_dir, args.output_dir)
    except (PreparationError, OSError, ValueError) as exc:
        print(f"GAGAL: {exc}", file=sys.stderr)
        return 1
    print(f"SUKSES: {manifest['source_file_count']} workbook Case 2 diproses.")
    for filename, count in manifest["record_counts"].items():
        print(f"  {filename}: {count}")
    print(f"Hasil: {args.output_dir.resolve()}")
    print("Belum diimpor ke Supabase; validasi RCA, energy, dan audit menyusul.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
